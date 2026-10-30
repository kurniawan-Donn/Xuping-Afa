"""
Membuat Buku Panduan (.docx) dari docs/manual/draft.md dengan python-docx.

    python docs/manual/scripts/build_docx.py

Setelah itu jalankan update_fields.ps1 agar Daftar Isi, Daftar Gambar, dan Daftar Tabel
terisi oleh Microsoft Word, sekaligus membuat PDF untuk pemeriksaan.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
DRAFT = ROOT / 'draft.md'
FONT = 'Arial'
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLACK = RGBColor(0, 0, 0)
TEXT_WIDTH_CM = 16.0          # A4 21 cm dikurangi margin 2 x 2,5 cm
IMG_WIDTH_CM = 15.0
IMG_MOBILE_WIDTH_CM = 6.2
IMG_MAX_HEIGHT_CM = 11.5      # agar gambar dan keterangannya muat di satu halaman
NOTE_FILL = {'Catatan': 'DEEAF6', 'Perhatian': 'FFF2CC'}


# ---------------------------------------------------------------- utilitas XML

def set_cell_fill(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=110, right=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for side, val in (('top', top), ('bottom', bottom), ('start', left), ('end', right)):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tc_pr.append(mar)


def set_table_borders(table, color='808080', size=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(size))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tbl_pr.append(borders)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    tr_pr.append(el)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement('w:cantSplit'))


def add_field(paragraph, instr, placeholder=''):
    """Field Word (PAGE, SEQ, TOC, ...) dengan nilai awal `placeholder`."""
    def fld(kind):
        el = OxmlElement('w:fldChar')
        el.set(qn('w:fldCharType'), kind)
        if kind == 'begin':
            el.set(qn('w:dirty'), 'true')
        return el

    run = paragraph.add_run()
    run._r.append(fld('begin'))
    run = paragraph.add_run()
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = f' {instr} '
    run._r.append(it)
    run = paragraph.add_run()
    run._r.append(fld('separate'))
    result = paragraph.add_run(placeholder)
    run = paragraph.add_run()
    run._r.append(fld('end'))
    return result


def bottom_border(paragraph, color='808080', size=4):
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    el = OxmlElement('w:bottom')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(size))
    el.set(qn('w:space'), '1')
    el.set(qn('w:color'), color)
    bdr.append(el)
    p_pr.append(bdr)


def picture_border(inline_shape, color='A6A6A6', width_emu=6350):
    """Garis tepi 0,5 pt di sekeliling gambar."""
    sp_pr = inline_shape._inline.graphic.graphicData.pic.spPr
    ln = OxmlElement('a:ln')
    ln.set('w', str(width_emu))
    fill = OxmlElement('a:solidFill')
    clr = OxmlElement('a:srgbClr')
    clr.set('val', color)
    fill.append(clr)
    ln.append(fill)
    sp_pr.append(ln)


def page_number_format(section, fmt, start=1):
    sect_pr = section._sectPr
    pg = sect_pr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType')
        sect_pr.append(pg)
    pg.set(qn('w:fmt'), fmt)
    pg.set(qn('w:start'), str(start))


# ---------------------------------------------------------------- gaya dokumen

def setup_styles(doc):
    styles = doc.styles

    def font(style, size, bold=False, color=BLACK, italic=False):
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = color
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.append(rfonts)
        for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rfonts.set(qn(attr), FONT)
        for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]

    normal = styles['Normal']
    font(normal, 11)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)

    for name, size, before, after in (('Heading 1', 14, 0, 18), ('Heading 2', 12, 12, 6), ('Heading 3', 11, 10, 4)):
        st = styles[name]
        font(st, size, bold=True, color=NAVY)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.15
    h1 = styles['Heading 1'].paragraph_format
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles['Heading 1'].font.all_caps = True

    cap = styles['Caption']
    font(cap, 10, bold=False, color=BLACK)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(12)

    # Judul bagian awal (Riwayat Revisi, Kata Pengantar, Daftar ...): tampil seperti bab
    # tetapi tidak ikut masuk Daftar Isi.
    front = styles.add_style('Judul Bagian', WD_STYLE_TYPE.PARAGRAPH)
    front.base_style = normal
    font(front, 14, bold=True, color=NAVY)
    front.font.all_caps = True
    front.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    front.paragraph_format.space_after = Pt(18)

    step = styles.add_style('Langkah', WD_STYLE_TYPE.PARAGRAPH)
    step.base_style = normal
    step.paragraph_format.left_indent = Cm(0.75)
    step.paragraph_format.first_line_indent = Cm(-0.75)
    step.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
    step.paragraph_format.space_after = Pt(4)

    cell = styles.add_style('Isi Tabel', WD_STYLE_TYPE.PARAGRAPH)
    cell.base_style = normal
    font(cell, 10)
    cell.paragraph_format.space_after = Pt(0)
    cell.paragraph_format.line_spacing = 1.0

    for level in (1, 2, 3):
        name = f'TOC {level}'
        st = styles[name] if name in [s.name for s in styles] else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        font(st, 11, bold=(level == 1))
        st.paragraph_format.left_indent = Cm(0.6 * (level - 1))
        st.paragraph_format.space_after = Pt(3)
    tof = styles['table of figures'] if 'table of figures' in [s.name for s in styles] else styles.add_style('table of figures', WD_STYLE_TYPE.PARAGRAPH)
    font(tof, 11)
    tof.paragraph_format.space_after = Pt(3)


def enable_update_fields(doc):
    settings = doc.settings.element
    el = OxmlElement('w:updateFields')
    el.set(qn('w:val'), 'true')
    settings.append(el)


# ---------------------------------------------------------------- teks inline

def smart_quotes(text):
    out, open_q = [], True
    for i, ch in enumerate(text):
        if ch == '"':
            prev = text[i - 1] if i else ' '
            open_q = prev in ' (\t\n/'
            out.append('“' if open_q else '”')
        else:
            out.append(ch)
    return ''.join(out)


def add_inline(paragraph, text, size=None, bold=False):
    text = smart_quotes(text)
    for i, part in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = bold or (i % 2 == 1)
        if size:
            run.font.size = Pt(size)


# ---------------------------------------------------------------- blok konten

class Builder:
    def __init__(self):
        self.doc = Document()
        self.chapter = 0
        self.fig = 0
        self.tab = 0
        self.figures = 0
        self.tables = 0
        self.pending_caption = None
        self.first_chapter = True
        setup_styles(self.doc)
        enable_update_fields(self.doc)
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
            setattr(sec, side, Cm(2.5))
        sec.header_distance = Cm(1.25)
        sec.footer_distance = Cm(1.25)

    # --- paragraf
    def para(self, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, style=None):
        p = self.doc.add_paragraph(style=style)
        p.alignment = align
        add_inline(p, text)
        return p

    def step(self, num, text):
        p = self.doc.add_paragraph(style='Langkah')
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run(f'{num}.\t')
        add_inline(p, text)

    def heading(self, level, text):
        if level == 1:
            m = re.match(r'BAB (\d+) (.+)', text)
            self.chapter = int(m.group(1))
            self.fig = self.tab = 0
            if self.first_chapter:
                self.start_main_section()
                self.first_chapter = False
            p = self.doc.add_paragraph(f'BAB {m.group(1)} {m.group(2).upper()}', style='Heading 1')
            if self.chapter > 1:
                p.paragraph_format.page_break_before = True
        else:
            self.doc.add_paragraph(text, style=f'Heading {level}')

    # --- keterangan gambar dan tabel memakai field SEQ agar Daftar Gambar/Tabel Word berfungsi
    def caption(self, kind, text):
        if kind == 'Gambar':
            self.fig += 1
            self.figures += 1
            num = self.fig
        else:
            self.tab += 1
            self.tables += 1
            num = self.tab
        p = self.doc.add_paragraph(style='Caption')
        p.add_run(f'{kind} {self.chapter}.')
        reset = f' \\r {num}' if num == 1 else ''
        add_field(p, f'SEQ {kind} \\* ARABIC{reset}', str(num))
        p.add_run(' ')
        add_inline(p, text)
        if kind == 'Tabel':
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
        return p

    def image(self, path, alt, mobile):
        full = ROOT / path
        w_px, h_px = Image.open(full).size
        width = IMG_MOBILE_WIDTH_CM if mobile else IMG_WIDTH_CM
        if width * h_px / w_px > IMG_MAX_HEIGHT_CM:
            width = IMG_MAX_HEIGHT_CM * w_px / h_px
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(0)
        shape = p.add_run().add_picture(str(full), width=Cm(width))
        picture_border(shape)
        self.caption('Gambar', alt)

    def table(self, rows, caption=None, widths=None, header=True):
        if caption:
            self.caption('Tabel', caption)
        ncol = len(rows[0])
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        set_table_borders(t)
        widths = widths or auto_widths(rows)
        for r, row in enumerate(rows):
            set_cant_split(t.rows[r])
            for c, value in enumerate(row):
                cell = t.cell(r, c)
                cell.width = Cm(widths[c])
                set_cell_margins(cell)
                p = cell.paragraphs[0]
                p.style = self.doc.styles['Isi Tabel']
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline(p, value, bold=(header and r == 0))
                if header and r == 0:
                    set_cell_fill(cell, 'D9D9D9')
        if header:
            set_repeat_header(t.rows[0])
        # Tabel pendek dijaga utuh dalam satu halaman; tabel panjang cukup menjaga
        # agar baris judul tidak tertinggal sendirian di bawah halaman.
        keep_rows = len(rows) - 1 if len(rows) <= 12 else 1
        for r in range(keep_rows):
            for cell in t.rows[r].cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def note(self, label, text):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t, color='BFBFBF')
        set_cant_split(t.rows[0])
        cell = t.cell(0, 0)
        cell.width = Cm(TEXT_WIDTH_CM)
        set_cell_fill(cell, NOTE_FILL[label])
        set_cell_margins(cell, 100, 100, 160, 160)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f'{label}: ')
        r.bold = True
        add_inline(p, text)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # --- bagian depan
    def cover(self, meta):
        def line(text, size, bold=False, before=0, after=0, color=BLACK):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(before)
            p.paragraph_format.space_after = Pt(after)
            r = p.add_run(text)
            r.font.size = Pt(size)
            r.bold = bold
            r.font.color.rgb = color
            return p

        line('BUKU PANDUAN PENGGUNAAN', 24, True, before=150, after=10, color=NAVY)
        line(meta['aplikasi'], 16, True, after=4)
        p = line(meta['url'], 11, after=0)
        bottom_border(p, color='1F3864', size=8)
        line('Disusun untuk', 11, before=110, after=2)
        line(meta['klien'], 14, True, after=40)
        line(f"Versi {meta['versi']}", 11, after=2)
        line(meta['tanggal'], 11, after=40)
        line('Disusun oleh', 11, after=2)
        line(meta['penyusun'], 12, True)

    def front_section(self):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        page_number_format(sec, 'lowerRoman', 1)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(fp, 'PAGE', 'i')
        for r in fp.runs:
            r.font.size = Pt(10)

    def front_title(self, text, new_page=True):
        p = self.doc.add_paragraph(text, style='Judul Bagian')
        p.paragraph_format.page_break_before = new_page

    def toc_page(self, title, instr, first=False):
        self.front_title(title, new_page=not first)
        p = self.doc.add_paragraph()
        add_field(p, instr, 'Klik kanan di sini lalu pilih Update Field untuk menampilkan daftar.')

    def start_main_section(self):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        page_number_format(sec, 'decimal', 1)
        hp = sec.header.paragraphs[0]
        # Style Header bawaan punya tab tengah di 8,25 cm; buang agar teks kanan benar-benar rata kanan.
        self.doc.styles['Header'].paragraph_format.tab_stops.clear_all()
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
        r = hp.add_run(f'{self.meta["aplikasi"]}\tBuku Panduan Penggunaan')
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        bottom_border(hp)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_field(fp, 'PAGE', '1')
        for r in fp.runs:
            r.font.size = Pt(10)

    def save(self, path):
        # Sampul (seksi pertama) tanpa header dan footer.
        first = self.doc.sections[0]
        first.header.is_linked_to_previous = False
        first.footer.is_linked_to_previous = False
        self.doc.save(path)


def auto_widths(rows):
    ncol = len(rows[0])
    if ncol == 3 and rows[0][-1] == 'Wajib':
        return [3.8, 9.2, 3.0]
    # Kolom berisi nilai pendek (Ya/Tidak) cukup selebar judulnya; sisanya dibagi rata
    # menurut panjang isi.
    fixed = {}
    for c in range(ncol):
        if max(len(r[c]) for r in rows[1:]) <= 6:
            fixed[c] = max(1.8, len(rows[0][c]) * 0.21 + 0.6)
    free = [c for c in range(ncol) if c not in fixed]
    lengths = {c: min(max(max(len(r[c]) for r in rows), 6), 60) for c in free}
    rest = TEXT_WIDTH_CM - sum(fixed.values())
    total = sum(lengths.values())
    return [fixed[c] if c in fixed else rest * lengths[c] / total for c in range(ncol)]


# ---------------------------------------------------------------- pembaca draf

def parse_table(lines, i):
    rows = []
    while i < len(lines) and lines[i].startswith('|'):
        cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
        if not all(re.fullmatch(r':?-+:?', c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def build():
    text = DRAFT.read_text(encoding='utf-8')
    text = re.sub(r'<!--.*?-->', '', text, flags=re.S)
    blocks = dict(re.findall(r'^::: (\w+)\n(.*?)^:::', text, flags=re.S | re.M))
    main = text[text.index('# BAB 1'):]

    b = Builder()
    meta = dict(line.split(': ', 1) for line in blocks['sampul'].strip().splitlines())
    b.meta = meta
    b.cover(meta)

    # Halaman awal bernomor romawi.
    b.front_section()
    b.front_title('Riwayat Revisi Dokumen', new_page=False)
    rev_lines = [l for l in blocks['revisi'].splitlines() if l.startswith('|')]
    rows, _ = parse_table(rev_lines, 0)
    b.table(rows, widths=[2.0, 3.6, 6.8, 3.6])
    b.front_title('Kata Pengantar')
    for para in blocks['pengantar'].strip().split('\n\n'):
        b.para(para.strip())
    pp = b.doc.add_paragraph()
    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pp.paragraph_format.space_before = Pt(24)
    pp.add_run(f"Madiun, {meta['tanggal']}\n\n\n{meta['penyusun'].split(' (')[0]}")
    b.toc_page('Daftar Isi', 'TOC \\o "1-3" \\h \\z \\u')
    b.toc_page('Daftar Gambar', 'TOC \\h \\z \\c "Gambar"')
    b.toc_page('Daftar Tabel', 'TOC \\h \\z \\c "Tabel"')

    lines = main.splitlines()
    i = 0
    para_buf = []

    def flush():
        if para_buf:
            b.para(' '.join(para_buf).strip())
            para_buf.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            flush()
            i += 1
            continue
        m = re.match(r'^(#{1,3}) (.+)', line)
        if m:
            flush()
            b.heading(len(m.group(1)), m.group(2))
            i += 1
            continue
        m = re.match(r'^!\[(.+?)\]\((.+?)\)(\{hp\})?$', line)
        if m:
            flush()
            b.image(m.group(2), m.group(1), bool(m.group(3)))
            i += 1
            continue
        m = re.match(r'^Tabel: (.+)', line)
        if m:
            flush()
            i += 1
            while not lines[i].strip():
                i += 1
            rows, i = parse_table(lines, i)
            b.table(rows, caption=m.group(1))
            continue
        m = re.match(r'^> \*\*(Catatan|Perhatian):\*\* (.+)', line)
        if m:
            flush()
            b.note(m.group(1), m.group(2))
            i += 1
            continue
        m = re.match(r'^(\d+)\. (.+)', line)
        if m:
            flush()
            b.step(m.group(1), m.group(2))
            i += 1
            continue
        para_buf.append(line.strip())
        i += 1
    flush()

    name = 'Buku_Panduan_' + meta['aplikasi'].replace(' ', '_') + f"_v{meta['versi']}.docx"
    out = ROOT / name
    b.save(out)
    print(f'{out} | {b.figures} gambar, {b.tables + 1} tabel')
    return out


if __name__ == '__main__':
    build()
